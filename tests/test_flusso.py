"""Flusso completo senza Claude: PDF inventati -> preparazione -> schede simulate -> salva -> componi."""

import json
import zipfile

from fascicolo_inventato import fascicolo, scheda
from openpyxl import load_workbook
from pdf_minimo import pdf_con_testo

from traccia import cli
from traccia import lavoro as lv


def _crea_lavoro(radice, nome="prova"):
    cartella = radice / "lavori" / nome
    ingresso = cartella / "ingresso"
    ingresso.mkdir(parents=True)
    schede = []
    with zipfile.ZipFile(ingresso / "fascicolo.zip", "w") as z:
        for file, _, s in fascicolo():
            z.writestr(f"documenti/{file}", pdf_con_testo((s["trascrizione"] or "pagina").splitlines()))
            schede.append(s)
    return cartella, schede


def _leggi_tutto(cartella, schede):
    id_lavoro = cartella.name
    assert cli.prepara_lavoro(id_lavoro) == 0
    pagine = lv.pagine(cartella)
    ordine = {file: s for (file, _, s), s in zip(fascicolo(), schede, strict=True)}
    for p in pagine:
        risposta = cartella / "risposte" / f"{p['pid']}.json"
        risposta.parent.mkdir(exist_ok=True)
        risposta.write_text(json.dumps(ordine[p["file"].rsplit("/", 1)[-1]]), encoding="utf-8")
        assert cli.salva(id_lavoro, p["pid"]) == 0
    return pagine


def test_flusso_completo(radice):
    cartella, schede = _crea_lavoro(radice)
    pagine = _leggi_tutto(cartella, schede)
    assert len(pagine) == len(schede)
    assert lv.stato(cartella)["fase"] == "composizione"
    assert lv.prossimo_lavoro() == {"lavoro": "prova", "componi": True}

    assert cli.componi("prova") == 0
    uscita = cartella / "uscita"
    assert lv.stato(cartella)["fase"] == "pronto"
    wb = load_workbook(uscita / "tracciabilita.xlsx")
    assert {"Tracciabilità", "Provenienza", "Richieste ai fornitori", "Documenti non usati", "Metodo"} <= set(
        wb.sheetnames
    )
    ws = wb["Tracciabilità"]
    link = [c.hyperlink.target for r in ws.iter_rows(min_row=2) for c in r if c.hyperlink]
    assert link and all((uscita / t).exists() for t in link)
    assert (uscita / "fascicolo_ordinato" / "INDICE.xlsx").exists()
    assert list((uscita / "fascicolo_ordinato" / "03_Certificati_3.1").glob("*.pdf"))

    # zona esatta dal testo nativo del PDF
    prove = json.loads((uscita / "prove.json").read_text(encoding="utf-8"))
    colata = next(f for f in prove["fonti"].values() if f["valore"] == "H1001")
    assert colata["esatta"] and colata["zona"]


def test_stesse_letture_stessa_uscita(radice):
    cartella, schede = _crea_lavoro(radice)
    _leggi_tutto(cartella, schede)
    cli.componi("prova")
    primo = (cartella / "uscita" / "risultato.json").read_text(encoding="utf-8")
    cli.componi("prova")
    assert (cartella / "uscita" / "risultato.json").read_text(encoding="utf-8") == primo


def test_scheda_sbagliata_rifiutata_con_errori(radice, capsys):
    cartella, _ = _crea_lavoro(radice)
    cli.prepara_lavoro("prova")
    capsys.readouterr()
    s = scheda("certificato_31")
    s["testata"]["numero"] = {"v": "12345", "cit": "numero inventato non sulla pagina", "zona": []}
    (cartella / "risposte").mkdir()
    (cartella / "risposte" / "p0001.json").write_text(json.dumps(s), encoding="utf-8")
    assert cli.salva("prova", "p0001") == 1
    uscita = json.loads(capsys.readouterr().out)
    assert any("citazione" in e for e in uscita["errori"])


def test_pagine_lette_non_si_rileggono(radice):
    cartella, schede = _crea_lavoro(radice)
    _leggi_tutto(cartella, schede)
    secondo = radice / "lavori" / "bis"
    (secondo / "ingresso").mkdir(parents=True)
    for f in (cartella / "ingresso").glob("*.zip"):
        (secondo / "ingresso" / f.name).write_bytes(f.read_bytes())
    cli.prepara_lavoro("bis")
    assert lv.stato(secondo)["da_leggere"] == 0  # stesse pagine: letture dall'archivio


def test_tentativi_poi_pagina_dichiarata_non_letta(radice):
    cartella, _ = _crea_lavoro(radice)
    cli.prepara_lavoro("prova")
    for _ in range(lv.MAX_TENTATIVI):
        assert lv.prossimo_lavoro()["lavoro"] == "prova"
    s = lv.stato(cartella)
    assert s["da_leggere"] == 0 and len(s["non_lette"]) == len(lv.pagine(cartella))
