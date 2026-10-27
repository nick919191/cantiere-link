"""Fascicolo ordinato: un PDF per documento, pagine raddrizzate, pagine bianche tolte, nomi chiari, ordine della
catena. Gli originali non si toccano: si crea una copia. Usa solo le letture già in archivio."""

import io
import re
from pathlib import Path

import pypdfium2 as pdfium
from openpyxl import Workbook
from PIL import Image

CARTELLE = {
    "packing_costruttore": "01_Costruttore",
    "dop_costruttore": "01_Costruttore",
    "ddt": "02_DDT",
    "certificato_31": "03_Certificati_3.1",
    "lettera_accompagnamento": "03_Certificati_3.1",
    "dop_ce_prodotto": "04_DoP_CE_prodotto",
    "da_escludere": "99_Non_usati",
    "illeggibile": "99_Non_usati",
}
TIPI = {
    "packing_costruttore": "Tabella_costruttore",
    "dop_costruttore": "DoP_costruttore",
    "ddt": "DDT",
    "certificato_31": "Certificato_3.1",
    "lettera_accompagnamento": "Lettera",
    "dop_ce_prodotto": "DoP_CE",
    "da_escludere": "Non_pertinente",
    "illeggibile": "Illeggibile",
}


def _pulito(testo: str, massimo: int = 40) -> str:
    t = re.sub(r"[^\w.-]+", "_", testo or "", flags=re.UNICODE).strip("_.")
    return t[:massimo]


def _documenti(pagine: list[dict], schede: dict[str, dict]) -> list[dict]:
    """Pagine consecutive dello stesso file, stessa famiglia e stesso numero (o pagina di seguito) = un documento."""
    documenti = []
    for p in pagine:
        s = schede.get(p["pid"])
        if not s or s.get("bianca"):
            continue
        t = s.get("testata", {})
        numero = next((t[k]["v"] for k in ("numero", "numero_dop") if k in t and t[k]["v"]), "")
        chi = next(
            (
                t[k]["v"]
                for k in ("emittente", "produttore", "fabbricante", "costruttore", "titolo", "commessa")
                if k in t and t[k]["v"]
            ),
            "",
        )
        data = t.get("data", {}).get("v", "")
        ultimo = documenti[-1] if documenti else None
        if (
            ultimo
            and ultimo["file"] == p["file"]
            and ultimo["famiglia"] == s["famiglia"]
            and (not numero or numero == ultimo["numero"] or s.get("continua_pagina_prima"))
        ):
            ultimo["pagine"].append((p, s.get("orientamento", 0)))
            continue
        documenti.append(
            {
                "famiglia": s["famiglia"],
                "numero": numero,
                "chi": chi,
                "data": data,
                "file": p["file"],
                "pagine": [(p, s.get("orientamento", 0))],
            }
        )
    return documenti


def crea(cartella_lavoro: Path, pagine: list[dict], schede: dict[str, dict], destinazione: Path) -> list[list]:
    destinazione.mkdir(parents=True, exist_ok=True)
    indice = []
    contatori: dict[str, int] = {}
    for doc in _documenti(pagine, schede):
        cartella = CARTELLE[doc["famiglia"]]
        contatori[cartella] = contatori.get(cartella, 0) + 1
        nome = (
            "_".join(
                x
                for x in (
                    f"{contatori[cartella]:03d}",
                    TIPI[doc["famiglia"]],
                    _pulito(doc["chi"], 30),
                    _pulito(doc["numero"], 24) if doc["numero"] else "",
                    _pulito(doc["data"].replace("/", "-"), 12) if doc["data"] else "",
                )
                if x
            )
            + ".pdf"
        )
        uscita = pdfium.PdfDocument.new()
        sorgenti = []  # aperte fino al salvataggio: le pagine importate le usano
        for p, gradi in doc["pagine"]:
            sorgente = cartella_lavoro / "ingresso" / p["file"]
            if sorgente.suffix.lower() == ".pdf":
                src = pdfium.PdfDocument(sorgente)
                uscita.import_pages(src, [p["pagina"] - 1])
            else:  # immagine: diventa una pagina PDF
                buf = io.BytesIO()
                with Image.open(cartella_lavoro / "pagine" / f"{p['pid']}.png") as img:
                    img.convert("RGB").save(buf, format="PDF", resolution=150)
                src = pdfium.PdfDocument(buf.getvalue())
                uscita.import_pages(src, [0])
            sorgenti.append(src)
            ultima = uscita[len(uscita) - 1]
            ultima.set_rotation((ultima.get_rotation() + gradi) % 360)
            ultima.close()
        (destinazione / cartella).mkdir(exist_ok=True)
        uscita.save(destinazione / cartella / nome)
        uscita.close()
        for s in sorgenti:
            s.close()
        indice.append([f"{cartella}/{nome}", doc["file"], ", ".join(str(p["pagina"]) for p, _ in doc["pagine"])])
    wb = Workbook()
    ws = wb.active
    ws.title = "Indice"
    ws.append(["NUOVO NOME", "FILE ORIGINALE", "PAGINE ORIGINALI"])
    for r in indice:
        ws.append(r)
    for col, w in zip("ABC", (70, 60, 20), strict=True):
        ws.column_dimensions[col].width = w
    wb.save(destinazione / "INDICE.xlsx")
    return indice
