"""Excel per il cliente: tracciabilità con un collegamento per cella alla prova, più i fogli di servizio."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .motore import COLONNE, VERSIONE

INTESTAZIONE = Font(bold=True, color="FFFFFF")
SFONDO = PatternFill("solid", fgColor="1D1D1F")
COLORI_STATO = {"PROVATO": "E7F6EA", "DEDOTTO": "E8F1FC", "DA CONFERMARE": "FFF4E5", "MANCANTE": "FDECEE"}
LINK = Font(color="0071E3", underline="single")
AVVISI_FISSI = [
    "Lettura e riconoscimento dei documenti sono automatici e possono contenere errori: verificare con gli originali.",
    "PROVATO indica un legame scritto sui documenti; non sostituisce la verifica di un tecnico.",
    "Pagine non lette o mancanti non significano che il documento non esista.",
]


def _testo(v) -> str:
    return "" if v is None else str(v)


def _foglio(wb, titolo, intestazione, righe, larghezze):
    ws = wb.create_sheet(titolo)
    ws.append(intestazione)
    for riga in righe:
        ws.append([_testo(v) for v in riga])
    for r in ws.iter_rows(min_row=2):
        for cella in r:
            cella.data_type = "s"
            cella.alignment = Alignment(wrap_text=True, vertical="top")
    for cella in ws[1]:
        cella.font, cella.fill = INTESTAZIONE, SFONDO
        cella.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(larghezze, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    ws.page_setup.orientation = "landscape"
    return ws


def scrivi(percorso: Path, risultato: dict, prove: dict, servizio: dict) -> None:
    """servizio: link (n, colonna) -> percorso relativo, documenti, non_usati, avvisi, metodo."""
    wb = Workbook()
    wb.remove(wb.active)
    voci = risultato["voci"]
    righe = [[v["colonne"][c]["v"] for c in COLONNE] + [v["stato"], " ".join(v["note"])] for v in voci]
    ws = _foglio(wb, "Tracciabilità", COLONNE + ["STATO", "NOTE"], righe, [26] * 12 + [16, 60])
    for i, v in enumerate(voci, start=2):
        for j, c in enumerate(COLONNE, start=1):
            link = servizio["link"].get((v["n"], c))
            if link and v["colonne"][c]["fonti"]:
                cella = ws.cell(row=i, column=j)
                cella.hyperlink = link
                cella.font = LINK
        ws.cell(row=i, column=13).fill = PatternFill("solid", fgColor=COLORI_STATO[v["stato"]])

    fonti = prove["fonti"]
    provenienza = []
    for v in voci:
        for c in COLONNE:
            for fid in v["colonne"][c]["fonti"]:
                f = fonti[fid]
                esatta = "esatta" if f["esatta"] else ("indicata da Claude" if f["zona"] else "")
                provenienza.append(
                    [
                        v["n"],
                        v["voce"],
                        c,
                        v["colonne"][c]["v"],
                        f["file"],
                        f["pagina"],
                        f["cit"],
                        ",".join(map(str, f["zona"])),
                        esatta,
                        servizio["link"].get((v["n"], c), ""),
                    ]
                )
    _foglio(
        wb,
        "Provenienza",
        ["N", "VOCE", "COLONNA", "VALORE", "DOCUMENTO", "PAGINA", "CITAZIONE", "ZONA 0-1000", "PRECISIONE", "PROVA"],
        provenienza,
        [6, 28, 22, 30, 40, 8, 50, 16, 16, 30],
    )
    _foglio(
        wb,
        "Prove",
        ["N", "VOCE", "STATO", "PAGINE USATE"],
        [[v["n"], v["voce"], v["stato"], servizio["pagine_voce"](v)] for v in voci],
        [6, 34, 16, 90],
    )
    _foglio(
        wb,
        "Motivi",
        ["N", "VOCE", "STATO", "MOTIVI", "AVVISI", "SCARTATI"],
        [
            [v["n"], v["voce"], v["stato"], "\n".join(v["motivi"]), "\n".join(v["avvisi"]), "\n".join(v["scartati"])]
            for v in voci
        ],
        [6, 30, 16, 60, 50, 50],
    )
    richieste = {}
    for v in voci:
        for r in v["richieste"]:
            richieste.setdefault((v["fornitore"] or "fornitore non individuato", r), []).append(v["voce"])
    _foglio(
        wb,
        "Richieste ai fornitori",
        ["FORNITORE", "DA CHIEDERE", "VOCI"],
        [[f, r, "; ".join(sorted(set(vs)))] for (f, r), vs in sorted(richieste.items())],
        [36, 50, 70],
    )
    _foglio(
        wb, "Documenti non usati", ["DOCUMENTO", "CONTENUTO", "PERCHÉ NON ENTRA"], servizio["non_usati"], [50, 50, 60]
    )
    _foglio(
        wb,
        "Documenti",
        ["FILE", "PAGINE", "LETTE", "NON LETTE", "FAMIGLIE"],
        servizio["documenti"],
        [50, 10, 10, 12, 60],
    )
    _foglio(wb, "Avvisi", ["AVVISO"], [[a] for a in servizio["avvisi"] + AVVISI_FISSI], [140])
    _foglio(wb, "Metodo", ["REGOLA"], [[m] for m in servizio["metodo"] + [f"Motore {VERSIONE}."]], [140])
    wb.save(percorso)
